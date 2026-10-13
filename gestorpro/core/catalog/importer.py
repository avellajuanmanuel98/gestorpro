"""
Importación de productos o ingredientes desde Excel (.xlsx) o CSV.

Flujo en dos pasos con el MISMO archivo:
  1. vista previa (`dry_run`): cada fila queda como nueva, ya existe o con errores;
     no se guarda nada;
  2. confirmación: se crean las filas válidas y nuevas en UNA transacción (si una
     falla, no se crea ninguna). Las que ya existen (mismo código o nombre) se
     omiten: la importación nunca modifica ítems existentes.

Cada fila pasa por el mismo serializer y las mismas reglas que el formulario
(services.normalize, permisos de costos, límite del plan), y la existencia
inicial entra como "saldo inicial" en el libro de inventario.

Números en formato colombiano: "1.500" = mil quinientos, "2,5" = dos y medio.
"""
import csv
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation

from django.db import transaction
from rest_framework.exceptions import ValidationError

from gestorpro.core import entitlements

from . import services
from .models import Category, Item, UnitOfMeasure

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 2000

COLUMNS = {
    'name': ('nombre', 'producto', 'ingrediente', 'insumo', 'descripcion del producto'),
    'code': ('codigo', 'sku', 'referencia', 'ref'),
    'category': ('categoria', 'grupo', 'linea'),
    'kind': ('tipo', 'tipo de producto'),
    'unit': ('unidad', 'se vende por', 'unidad de medida', 'medida'),
    'price': ('precio', 'precio sin iva', 'precio de venta', 'precio venta', 'valor'),
    'tax_rate': ('iva', 'iva %', 'impuesto', '% iva'),
    'avg_cost': ('costo', 'costo unitario', 'costo por unidad', 'precio de compra'),
    'stock': ('existencia', 'existencia inicial', 'cantidad', 'stock', 'inventario'),
    'minimum_stock': ('minimo', 'existencia minima', 'stock minimo'),
}
KINDS = {
    'elaborado': Item.Kind.FINISHED_GOOD, 'produccion': Item.Kind.FINISHED_GOOD, 'propio': Item.Kind.FINISHED_GOOD,
    'reventa': Item.Kind.RESALE, 'comprado': Item.Kind.RESALE, 'servicio': Item.Kind.SERVICE,
}
TEMPLATE = {
    'products': ['Nombre', 'Código', 'Categoría', 'Tipo', 'Unidad', 'Precio sin IVA', 'IVA %', 'Costo',
                 'Existencia inicial', 'Existencia mínima'],
    'ingredients': ['Nombre', 'Código', 'Categoría', 'Unidad', 'Costo', 'Existencia inicial', 'Existencia mínima'],
}
EXAMPLES = {
    ('products', True): [['Pan de bono', '', 'Panes', 'Elaborado', 'und', '1200', '0', '450', '40', '10'],
                         ['Gaseosa 400 ml', '', 'Bebidas frías', 'Reventa', 'und', '2521', '19', '1700', '24', '6']],
    ('products', False): [['Producto de ejemplo', '', 'General', 'Reventa', 'und', '10000', '19', '6000', '12', '3'],
                          ['Servicio de ejemplo', '', 'Servicios', 'Servicio', 'und', '25000', '19', '', '', '']],
    ('ingredients', True): [['Harina de trigo', '', 'Harinas y almidones', 'kg', '3200', '50', '20'],
                            ['Huevos', '', 'Lácteos y huevos', 'und', '450', '180', '60']],
    ('ingredients', False): [['Insumo de ejemplo', '', 'Materias primas', 'kg', '5000', '10', '2'],
                             ['Empaque de ejemplo', '', 'Empaques', 'und', '300', '200', '50']],
}


def _key(text) -> str:
    text = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9% ]', ' ', text)).strip()


HEADER_MAP = {_key(alias): field for field, aliases in COLUMNS.items() for alias in aliases}


def number(value) -> Decimal | None:
    """Número de una celda: acepta 1500, '1.500', '$ 1.500', '2,5', '19%'. None si está vacía."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return Decimal(str(value))
    text = re.sub(r'[\s$%]', '', str(value))
    if ',' in text and '.' in text:  # el último separador es el decimal
        text = text.replace('.', '').replace(',', '.') if text.rfind(',') > text.rfind('.') else text.replace(',', '')
    elif ',' in text:
        text = text.replace(',', '.')
    elif re.fullmatch(r'-?\d{1,3}(\.\d{3})+', text):  # 1.500 o 1.234.567: separador de miles
        text = text.replace('.', '')
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f'"{value}" no es un número') from exc


def read_rows(upload) -> list[dict]:
    """Lee el archivo y devuelve filas {campo: valor} con el número de línea del archivo."""
    if upload.size > MAX_BYTES:
        raise ValidationError({'file': 'El archivo pesa más de 2 MB. Divídelo en partes.'})
    name = upload.name.lower()
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            sheet = load_workbook(upload, read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise ValidationError({'file': 'No se pudo leer el Excel. Guárdalo como .xlsx y reintenta.'}) from exc
        table = [list(r) for r in sheet.iter_rows(values_only=True)]
    elif name.endswith('.csv'):
        raw = upload.read()
        for encoding in ('utf-8-sig', 'cp1252'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        delimiter = ';' if text.split('\n', 1)[0].count(';') >= text.split('\n', 1)[0].count(',') else ','
        table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    else:
        raise ValidationError({'file': 'Sube un archivo de Excel (.xlsx) o CSV.'})

    while table and not any(str(c or '').strip() for c in table[0]):
        table.pop(0)
    if not table:
        raise ValidationError({'file': 'El archivo está vacío.'})
    fields = [HEADER_MAP.get(_key(h)) for h in table[0]]
    if 'name' not in fields:
        raise ValidationError({'file': 'Falta la columna "Nombre". Usa la plantilla para ver las columnas.'})
    rows = []
    for line, values in enumerate(table[1:], start=2):
        if not any(str(v or '').strip() for v in values):
            continue
        rows.append({'_line': line, **{f: v for f, v in zip(fields, values, strict=False) if f}})
    if len(rows) > MAX_ROWS:
        raise ValidationError({'file': f'Máximo {MAX_ROWS} filas por archivo.'})
    if not rows:
        raise ValidationError({'file': 'El archivo no tiene filas con datos.'})
    return rows


class Importer:
    def __init__(self, group: str, request):
        if group not in ('products', 'ingredients'):
            raise ValidationError({'group': 'Indica products o ingredients.'})
        self.group, self.request = group, request
        self.can_costs = request.membership.has_perm('catalog.view_costs')
        self.category_kind = Category.Kind.PRODUCT if group == 'products' else Category.Kind.INGREDIENT
        self.categories = {_key(c.name): c for c in Category.objects.filter(kind=self.category_kind)}
        self.units = {}
        for u in UnitOfMeasure.objects.all():
            for alias in (u.code, u.symbol, u.name):
                self.units.setdefault(_key(alias), u)
        self.units.update({k: self.units[v] for k, v in (('unidad', 'und'), ('unidades', 'und'), ('u', 'und'),
                                                         ('kilo', 'kg'), ('kilos', 'kg'), ('gramos', 'g'),
                                                         ('litros', 'l'), ('lt', 'l')) if v in self.units})
        existing = Item.objects.values_list('code', 'name')
        self.codes = {_key(c) for c, _ in existing}
        self.names = {_key(n) for _, n in existing}

    def _payload(self, row: dict) -> tuple[dict, list[str], str | None]:
        """Datos para ItemSerializer, errores de formato y nombre de una categoría a crear."""
        errors, new_category = [], None
        data = {'name': str(row.get('name') or '').strip()[:200]}
        if row.get('code') not in (None, ''):
            data['code'] = str(row['code']).strip()[:50]
        if self.group == 'ingredients':
            data['kind'] = Item.Kind.RAW_MATERIAL
        else:
            kind = _key(row.get('kind'))
            data['kind'] = next((v for k, v in KINDS.items() if k in kind), None) if kind \
                else Item.Kind.FINISHED_GOOD
            if data['kind'] is None:
                errors.append(f'Tipo "{row.get("kind")}" no reconocido: usa Elaborado, Reventa o Servicio.')
        unit = _key(row.get('unit'))
        if unit:
            if unit in self.units:
                data['unit'] = self.units[unit].code
            else:
                errors.append(f'Unidad "{row.get("unit")}" no reconocida (ej.: und, kg, g, l, ml, lb).')
        category = str(row.get('category') or '').strip()
        if category:
            found = self.categories.get(_key(category))
            if found:
                data['category'] = found.id
            else:
                new_category = category[:100]
        for field in ('price', 'tax_rate', 'avg_cost', 'stock', 'minimum_stock'):
            if field == 'avg_cost' and not self.can_costs:
                continue
            try:
                value = number(row.get(field))
            except ValueError as exc:
                errors.append(f'{field}: {exc}')
                continue
            if value is None:
                continue
            if field == 'tax_rate' and value < 1 and value != 0:
                value *= 100  # 0,19 → 19 %
            if value < 0:
                errors.append(f'{field}: no puede ser negativo')
                continue
            data[field] = value
        return data, errors, new_category

    def evaluate(self) -> list[dict]:
        rows = read_rows(self.request.FILES.get('file') or _missing())
        from .serializers import ItemSerializer

        results, seen = [], set()
        for row in rows:
            data, errors, new_category = self._payload(row)
            name_key, code_key = _key(data['name']), _key(data.get('code'))
            status = 'new'
            if not data['name']:
                errors.append('Falta el nombre.')
            elif name_key in self.names or (code_key and code_key in self.codes):
                status = 'exists'
            elif name_key in seen:
                errors.append('Nombre repetido en el archivo.')
            seen.add(name_key)
            if status == 'new' and not errors:
                serializer = ItemSerializer(data=data, context={'request': self.request})
                if not serializer.is_valid():
                    errors += [f'{k}: {" ".join(map(str, v))}' for k, v in serializer.errors.items()]
            results.append({'line': row['_line'], 'name': data['name'], 'code': data.get('code', ''),
                            'kind': data.get('kind'), 'unit': data.get('unit', 'und'),
                            'price': str(data['price']) if 'price' in data else None,
                            'stock': str(data['stock']) if 'stock' in data else None,
                            'new_category': new_category,
                            'status': 'error' if errors else status, 'errors': errors, '_data': data})
        return results

    @transaction.atomic
    def commit(self, results: list[dict]) -> list[Item]:
        from .serializers import ItemSerializer
        from .views import create_item

        new = [r for r in results if r['status'] == 'new']
        products = sum(1 for r in new if r['kind'] in Item.PRODUCT_KINDS)
        if products:
            entitlements.check_limit(self.request.membership.tenant_id, 'products', services.products_count,
                                     adding=products)
        created = []
        for r in new:
            data = dict(r['_data'])
            if r['new_category']:
                key = _key(r['new_category'])
                if key not in self.categories:
                    self.categories[key] = Category.objects.create(name=r['new_category'], kind=self.category_kind)
                data['category'] = self.categories[key].id
            serializer = ItemSerializer(data=data, context={'request': self.request})
            serializer.is_valid(raise_exception=True)
            created.append(create_item(serializer, self.request))
        return created


def _missing():
    raise ValidationError({'file': 'Adjunta el archivo.'})


def template(group: str, bakery: bool) -> bytes:
    """Plantilla .xlsx con los encabezados y dos filas de ejemplo según el tipo de negocio."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Productos' if group == 'products' else 'Ingredientes'
    ws.append(TEMPLATE[group])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in EXAMPLES[(group, bakery)]:
        ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 4)
    help_ws = wb.create_sheet('Instrucciones')
    for line in (
        'Borra las filas de ejemplo y escribe una fila por ítem. Solo "Nombre" es obligatorio.',
        'Código: déjalo vacío y se asigna solo.',
        'Tipo (productos): Elaborado, Reventa o Servicio. Vacío = Elaborado.',
        'Unidad: und, kg, g, l, ml, lb (o el nombre: unidad, kilo, litro). Vacío = und.',
        'Precio: SIN IVA. El sistema suma el IVA al vender. IVA: 0, 5 o 19.',
        'Costo y existencia: la existencia entra como saldo inicial al costo indicado.',
        'Si un ítem ya existe (mismo nombre o código), se omite: la importación no modifica lo existente.',
        'Las categorías que no existan se crean solas.',
    ):
        help_ws.append([line])
    help_ws.column_dimensions['A'].width = 100
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
