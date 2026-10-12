"""Importación de productos e ingredientes desde Excel o CSV."""
import io
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from gestorpro.core.catalog.importer import number
from gestorpro.core.catalog.models import Category, Item
from gestorpro.core.inventory.models import StockMovement
from gestorpro.core.tenancy.context import tenant_context
from tests import support as s

pytestmark = pytest.mark.django_db
URL = '/api/catalog/import/'


def xlsx(rows):
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile('productos.xlsx', out.getvalue())


def csv_file(text, encoding='utf-8-sig'):
    return SimpleUploadedFile('productos.csv', text.encode(encoding))


@pytest.fixture
def shop():
    tenant = s.make_tenant('Importadora')
    return {'t': tenant, 'admin': s.api_for(s.make_user(tenant, 'adm@imp.co', 'ADMIN'))}


def test_colombian_numbers_are_read_correctly():
    assert number('1.500') == Decimal('1500') and number('$ 1.234.567') == Decimal('1234567')
    assert number('2,5') == Decimal('2.5') and number('1.500,75') == Decimal('1500.75')
    assert number('19%') == Decimal('19') and number(1200) == Decimal('1200') and number('') is None
    with pytest.raises(ValueError):
        number('mil')


def test_preview_saves_nothing_and_flags_each_row(shop):
    s.make_product(shop['t'], code='PAN-1')  # "Producto Importadora" ya existe
    upload = xlsx([
        ['Nombre', 'Código', 'Categoría', 'Tipo', 'Unidad', 'Precio sin IVA', 'IVA %', 'Costo', 'Existencia inicial'],
        ['Pan de bono', '', 'Panes', 'Elaborado', 'und', '1.200', 0, 450, 40],
        ['Producto Importadora', '', '', 'Reventa', 'und', 1000, 0, '', ''],
        ['Gaseosa', '', 'Bebidas', 'Volador', 'und', 2521, 19, 1700, 24],
        ['Torta', '', '', '', 'bandeja', 5000, 0, '', ''],
        ['Sin precio', '', '', 'Reventa', 'und', '', 0, '', ''],
        ['', '', '', '', '', '', '', '', ''],
    ])
    res = shop['admin'].post(URL + '?group=products', {'file': upload}, format='multipart')
    assert res.status_code == 200
    status = {r['name']: r['status'] for r in res.data['rows']}
    assert status == {'Pan de bono': 'new', 'Producto Importadora': 'exists', 'Gaseosa': 'error',
                      'Torta': 'error', 'Sin precio': 'error'}
    assert res.data['summary'] == {'new': 1, 'exists': 1, 'error': 3}
    assert res.data['new_categories'] == ['Panes']
    with tenant_context(shop['t']):
        assert not Item.objects.filter(name='Pan de bono').exists()


def test_confirm_creates_items_categories_and_opening_stock(shop):
    upload = csv_file('Nombre;Categoría;Tipo;Precio sin IVA;IVA %;Costo;Existencia inicial;Existencia mínima\n'
                      'Pan de bono;Panes;Elaborado;1.200;0;450;40;10\n'
                      'Gaseosa 400 ml;Bebidas;Reventa;2521;19;1.700;24;6\n')
    res = shop['admin'].post(URL + '?group=products&confirm=1', {'file': upload}, format='multipart')
    assert res.status_code == 200 and res.data['created'] == 2
    with tenant_context(shop['t']):
        pan = Item.objects.get(name='Pan de bono')
        assert (pan.kind, pan.price, pan.avg_cost, pan.stock, pan.minimum_stock) == (
            Item.Kind.FINISHED_GOOD, Decimal('1200'), Decimal('450'), Decimal('40'), Decimal('10'))
        assert pan.code and pan.category.name == 'Panes' and pan.category.kind == Category.Kind.PRODUCT
        opening = StockMovement.objects.get(item=pan)
        assert opening.type == StockMovement.Type.OPENING and opening.quantity == Decimal('40')
    # Reimportar el mismo archivo no duplica nada
    upload = csv_file('Nombre;Precio sin IVA\nPan de bono;1200\n')
    again = shop['admin'].post(URL + '?group=products&confirm=1', {'file': upload}, format='multipart')
    assert again.data['created'] == 0 and again.data['exists'] == 1


def test_ingredients_use_their_unit_and_never_become_sellable(shop):
    upload = csv_file('Nombre,Unidad,Costo,Existencia inicial\nHarina de trigo,kilo,3200,50\nLeche,litro,"4.200",12\n')
    res = shop['admin'].post(URL + '?group=ingredients&confirm=1', {'file': upload}, format='multipart')
    assert res.data['created'] == 2
    with tenant_context(shop['t']):
        flour = Item.objects.get(name='Harina de trigo')
        assert (flour.kind, flour.unit.code, flour.is_sellable, flour.stock) == (
            Item.Kind.RAW_MATERIAL, 'kg', False, Decimal('50'))


def test_costs_are_ignored_for_who_cannot_see_them(shop):
    tenant = shop['t']
    api = s.api_for(s.make_user(tenant, 'inv@imp.co', 'INVENTORY'))
    from gestorpro.core.access.models import Role
    with tenant_context(tenant):
        Role.objects.get(code='INVENTORY').permissions.remove('catalog.view_costs')
    upload = csv_file('Nombre;Unidad;Costo\nAzúcar;kg;99999\n')
    res = api.post(URL + '?group=ingredients&confirm=1', {'file': upload}, format='multipart')
    assert res.status_code == 200 and res.data['created'] == 1
    with tenant_context(tenant):
        assert Item.objects.get(name='Azúcar').avg_cost == 0


def test_plan_limit_is_respected_for_the_whole_file(shop):
    tenant = s.make_tenant('Pequeña', plan_code='starter')  # 150 productos
    api = s.api_for(s.make_user(tenant, 'adm@peq.co', 'ADMIN'))
    rows = 'Nombre;Precio sin IVA\n' + ''.join(f'Producto {i};1000\n' for i in range(151))
    res = api.post(URL + '?group=products&confirm=1', {'file': csv_file(rows)}, format='multipart')
    assert res.status_code == 403
    with tenant_context(tenant):
        assert not Item.objects.filter(name__startswith='Producto ').exists()


def test_template_and_permissions(shop):
    res = shop['admin'].get(URL, {'group': 'products'})
    assert res.status_code == 200
    header = [c.value for c in next(load_workbook(io.BytesIO(res.content)).active.iter_rows(max_row=1))]
    assert header[0] == 'Nombre' and 'Precio sin IVA' in header
    cashier = s.api_for(s.make_user(shop['t'], 'caj@imp.co', 'CASHIER'))
    denied = cashier.post(URL + '?group=products', {'file': csv_file('Nombre\nX\n')}, format='multipart')
    assert denied.status_code == 403
    bad = shop['admin'].post(URL + '?group=products', {'file': SimpleUploadedFile('x.pdf', b'%PDF')},
                             format='multipart')
    assert bad.status_code == 400
